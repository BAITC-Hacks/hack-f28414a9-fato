from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any
import os

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape


def create_report(data: dict[str, Any], fmt: str) -> Path:
    path = Path(tempfile.gettempdir()) / f"autoprotocol-{data['id']}.{fmt}"
    if fmt == "docx":
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(.65)
        section.bottom_margin = Inches(.65)
        doc.add_heading("Протокол совещания", 0)
        doc.add_paragraph(f"Запись: {data['filename']}  •  Язык: {data['language']}")
        doc.add_heading("Краткое содержание", level=1)
        doc.add_paragraph(data["summary"])
        doc.add_heading("Поручения", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, title in zip(table.rows[0].cells, ["Задача / исходная реплика", "Ответственный", "Срок", "Говорящий"]):
            cell.text = title
        for item in data["tasks"]:
            cells = table.add_row().cells
            cells[0].text = f"{item['task']}\nРеплика: {item['source']}"
            cells[1].text = item["responsible"]
            cells[2].text = item["deadline"]
            cells[3].text = item["speaker"]
        if not data["tasks"]:
            doc.add_paragraph("Явные поручения по ключевым словам не найдены.")
        doc.add_heading("Транскрипт", level=1)
        for item in data["transcript"]:
            p = doc.add_paragraph()
            run = p.add_run(f"{item['speaker']}  [{item['start']:.1f} с]: ")
            run.bold = True
            p.add_run(item["text"])
        doc.save(path)
    else:
        # Register a Unicode font when present so Cyrillic/Kazakh text renders in PDF.
        font_candidates = [Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                           Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf"]
        font_path = next((candidate for candidate in font_candidates if candidate.exists()), None)
        font_name = "Helvetica"
        if font_path:
            pdfmetrics.registerFont(TTFont("AutoProtocolUnicode", str(font_path)))
            font_name = "AutoProtocolUnicode"
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="Cyrillic", parent=styles["BodyText"], fontName=font_name, fontSize=8, leading=11))
        styles["Title"].fontName = font_name
        styles["Heading2"].fontName = font_name
        story = [Paragraph("Протокол совещания", styles["Title"]), Spacer(1, 5 * mm),
                 Paragraph(f"Запись: {escape(data['filename'])} &nbsp; Язык: {escape(data['language'])}", styles["Cyrillic"]),
                 Spacer(1, 4 * mm), Paragraph("Краткое содержание", styles["Heading2"]),
                 Paragraph(escape(data["summary"]), styles["Cyrillic"]), Spacer(1, 4 * mm),
                 Paragraph("Поручения", styles["Heading2"])]
        rows = [["Задача / исходная реплика", "Ответственный", "Срок", "Говорящий"]]
        rows.extend([[Paragraph(escape(f"{t['task']}\nРеплика: {t['source']}"), styles["Cyrillic"]),
                      Paragraph(escape(t["responsible"]), styles["Cyrillic"]),
                      Paragraph(escape(t["deadline"]), styles["Cyrillic"]),
                      Paragraph(escape(t["speaker"]), styles["Cyrillic"])] for t in data["tasks"]])
        if len(rows) == 1:
            rows.append(["Явные поручения не найдены", "—", "—", "—"])
        table = Table(rows, colWidths=[80 * mm, 35 * mm, 30 * mm, 30 * mm], repeatRows=1)
        table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font_name),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173b63")),
                                   ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                   ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d6deea")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 5),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                                   ("TOPPADDING", (0, 0), (-1, -1), 6),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        story.extend([table, Spacer(1, 5 * mm), Paragraph("Транскрипт", styles["Heading2"])])
        for item in data["transcript"]:
            story.append(Paragraph(f"<b>{escape(item['speaker'])} [{item['start']:.1f} с]</b>: {escape(item['text'])}", styles["Cyrillic"]))
            story.append(Spacer(1, 2 * mm))
        SimpleDocTemplate(str(path), pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm,
                          topMargin=14 * mm, bottomMargin=14 * mm).build(story)
    return path
